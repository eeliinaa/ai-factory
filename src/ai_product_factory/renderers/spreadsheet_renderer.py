from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from ..domain import SpreadsheetArtifactSpec
from ..outputs.file_renderer import normalize_artifact_file_name


def render_spreadsheet_artifact(output_dir: Path, file_name: str, spec: SpreadsheetArtifactSpec) -> tuple[Path, str]:
    normalized_name = normalize_artifact_file_name(file_name)
    file_path = output_dir / normalized_name
    file_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    for index, sheet_spec in enumerate(spec.sheets):
        worksheet = workbook.create_sheet(title=sheet_spec.name[:31] or f"Sheet{index + 1}")
        for column_index, column_spec in enumerate(sheet_spec.columns, start=1):
            cell = worksheet.cell(row=1, column=column_index, value=column_spec.header)
            cell.font = Font(bold=True)
            if column_spec.width:
                worksheet.column_dimensions[cell.column_letter].width = column_spec.width
        for row_index, row_values in enumerate(sheet_spec.rows, start=2):
            for column_index, value in enumerate(row_values, start=1):
                worksheet.cell(row=row_index, column=column_index, value=value)
        worksheet.freeze_panes = "A2"

    if not spec.sheets:
        worksheet = workbook.create_sheet(title="Sheet1")
        worksheet["A1"] = spec.workbook_title
        worksheet.freeze_panes = "A2"

    workbook.save(file_path)
    return file_path, "xlsx"